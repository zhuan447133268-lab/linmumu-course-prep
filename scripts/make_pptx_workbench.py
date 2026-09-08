# -*- coding: utf-8 -*-
"""02_课件.pptx —— 打造教师个人工作台（linmumu-course-prep 单节实测，工具/流程类课实例）
换主题 = 复制本文件，改内容。设计系统来自 deck_lib.py。
用法：python make_pptx_workbench.py [学科族] [输出pptx路径] [配图目录]
  默认：文科  ./02_课件.pptx  ./_figs   （参数顺序与 make_pptx_v2.py 一致）
  注意：配图先由 make_figs_workbench.py 生成到配图目录，再跑本脚本。
依赖：python-pptx。
"""
import os
import sys
from pptx.dml.color import RGBColor
import deck_lib as D

THEME = sys.argv[1] if len(sys.argv) > 1 else '文科'
OUT = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.path.join(os.getcwd(), '02_课件.pptx')
FIGS = os.path.abspath(sys.argv[3]) if len(sys.argv) > 3 else os.path.join(os.getcwd(), '_figs')
os.makedirs(os.path.dirname(OUT) or '.', exist_ok=True)

D.new_prs(theme=THEME, course='教师数字素养培训 · 第 1 节 打造教师个人工作台', total=14)
prs = D.prs
W, H = D.W, D.H
INK, ORANGE, ORANGE_SOFT = D.INK, D.ORANGE, D.ORANGE_SOFT
TEAL, TEAL_SOFT, BLUE = D.TEAL, D.TEAL_SOFT, D.BLUE
CREAM, WHITE, LINE, GRAY, FAINT = D.CREAM, D.WHITE, D.LINE, D.GRAY, D.FAINT
INK_SOFT = D.INK_SOFT
new_page, section_page, card, text, rect, oval = D.new_page, D.section_page, D.card, D.text, D.rect, D.oval
header, conclusion_bar, tag, demo_link, pic = D.header, D.conclusion_bar, D.tag, D.demo_link, D.pic

# ============ P1 封面 ============
D._n[0] += 1
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, W, H, INK)
oval(s, 8.6, 1.2, 6.6, 6.6, None, INK_SOFT, 2.4)
oval(s, 9.7, 2.3, 4.4, 4.4, None, INK_SOFT, 2.0)
oval(s, 10.6, 3.2, 2.6, 2.6, None, ORANGE, 1.8)
text(s, 0.95, 1.72, 8.0, 0.4, [('教师数字素养培训 · 第 1 节', 15, True, ORANGE)])
text(s, 0.95, 2.25, 9.6, 2.4, [
    ('打造教师', 46, True, WHITE),
    ('个人工作台', 46, True, WHITE),
], spacing=1.12)
rect(s, 0.98, 4.62, 1.55, 0.05, ORANGE)
text(s, 0.95, 4.88, 8.6, 0.5, [('一句话生成、双击即开、内容存自己电脑', 19, False, FAINT)])
text(s, 0.95, 6.35, 10.0, 0.4, [('面向高校 / 中小学教师 · 教师发展培训 · 2 学时', 13.5, False,
       RGBColor(0x9A, 0x8C, 0x7C))])
D.footer_dark(s)

# ============ P2 今天的问题 ============
s = new_page('今天的问题', '你的待办，散在多少个地方？')
items = [
    ('微信', '教研室群 @你、家长留言、学生提问', ORANGE),
    ('浏览器', '收藏了 20 篇文献，再没打开过', BLUE),
    ('便签 / 脑子', '「明天交材料」——然后忘了', TEAL),
    ('各类系统', '教务、科研、OA 各登一遍', INK_SOFT),
]
cw, gap, x0 = 2.85, 0.33, 0.9
for i, (t, d, c) in enumerate(items):
    x = x0 + i * (cw + gap)
    card(s, x, 2.0, cw, 2.7, WHITE, LINE)
    rect(s, x, 2.0, cw, 0.14, c, None, shape=D.MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x + 0.28, 2.45, cw - 0.5, 0.7, [(t, 22, True, c)])
    text(s, x + 0.28, 3.25, cw - 0.5, 1.3, [(d, 14.5, False, GRAY)], spacing=1.3)
conclusion_bar(s, [('东西越散，越记不住、越做不完。', 17, True, WHITE)], y=5.35, h=0.66)

# ============ P3 路线图 ============
s = new_page('本节路线', '看成品 → 跟做 → 带走，90 分钟')
rows = [
    ('01', '看成品', '逛一圈做好的工作台，认四块拼图', BLUE),
    ('02', '跟做', '一句话生成你的台子，加自己的模块', ORANGE),
    ('03', '带走', '迁移卡四步，课下自己接着搭', TEAL),
]
cw, gap, x0 = 3.72, 0.35, 0.9
for i, (num, t, d, c) in enumerate(rows):
    x = x0 + i * (cw + gap)
    card(s, x, 2.15, cw, 3.3, WHITE, LINE)
    rect(s, x, 2.15, cw, 0.14, c, None, shape=D.MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x + 0.35, 2.62, cw - 0.7, 0.85, [(num, 40, True, c)])
    text(s, x + 0.35, 3.62, cw - 0.7, 0.55, [(t, 21, True, INK)])
    text(s, x + 0.35, 4.28, cw - 0.7, 1.0, [(d, 14.5, False, GRAY)], spacing=1.25)

# ============ P4 章节 01 ============
section_page('01', '看成品', '先逛一圈，再决定你要哪几块')

# ============ P5 四块拼图 ============
s = new_page('第一站 · 拼图', '所有工作台，都是这四块里挑')
blocks = [
    ('板', '任务板', '事按「重要 / 着急」分四块', BLUE),
    ('钟', '专注钟', '干 25 / 50 / 90 分钟，响了就歇', ORANGE),
    ('窗', '资料窗', '翻自己的收藏 / 文献，只看不改', TEAL),
    ('搭', 'AI 搭子', '把今天的事整理成话，发 AI', INK_SOFT),
]
cw, gap, x0 = 2.85, 0.33, 0.9
for i, (g, t, d, c) in enumerate(blocks):
    x = x0 + i * (cw + gap)
    card(s, x, 2.1, cw, 3.2, WHITE, LINE)
    oval(s, x + cw / 2 - 0.5, 2.45, 1.0, 1.0, c, None)
    text(s, x + cw / 2 - 0.5, 2.45, 1.0, 1.0, [(g, 34, True, WHITE)],
         anchor=D.MSO_ANCHOR.MIDDLE, align=D.PP_ALIGN.CENTER)
    text(s, x + 0.28, 3.65, cw - 0.5, 0.55, [(t, 20, True, INK)], align=D.PP_ALIGN.CENTER)
    text(s, x + 0.28, 4.25, cw - 0.5, 0.95, [(d, 14, False, GRAY)],
         align=D.PP_ALIGN.CENTER, spacing=1.25)
conclusion_bar(s, [('不用四块全要——挑你最卡的 1~2 块就够。', 16.5, True, WHITE)], y=5.65, h=0.62)

# ============ P6 任务板 ============
s = new_page('第一站 · 任务板', '先把「重要」和「着急」分开')
quad = [
    ('重要且着急', '今天必做', ORANGE),
    ('重要不着急', '排期做', BLUE),
    ('不重要但着急', '速办 / 委托', GRAY),
    ('不重要不着急', '别做', TEAL),
]
qx, qy, qw, qh, gap = 0.95, 2.0, 5.45, 2.35, 0.25
for i, (t, d, c) in enumerate(quad):
    r, cc = divmod(i, 2)
    x = qx + cc * (qw + gap); y = qy + r * (qh + gap)
    card(s, x, y, qw, qh, WHITE, LINE)
    rect(s, x, y, 0.14, qh, c)
    text(s, x + 0.4, y + 0.3, qw - 0.6, 0.6, [(t, 19, True, c)])
    text(s, x + 0.4, y + 1.05, qw - 0.6, 0.8, [(d, 15, False, GRAY)])
card(s, 7.1, 2.0, 5.3, 4.7, ORANGE_SOFT, None)
text(s, 7.5, 2.35, 4.6, 4.0, [
    [('为什么是这块？', 18, True, ORANGE)],
    ('人每天被「着急但不重要」的事推着走，', 15, False, INK),
    ('真正重要的事永远往后拖。', 15, False, INK),
    ('', 8, False, INK),
    [('四块一分，', 16, True, ORANGE), ('先填「重要且着急」，', 16, False, INK)],
    ('剩下的再排期。', 16, False, INK),
], spacing=1.3, space_after=8)

# ============ P7 专注钟 ============
s = new_page('第一站 · 专注钟', '干 25 分钟，比干一下午更出活')
card(s, 0.9, 2.1, 5.55, 4.5, WHITE, LINE)
oval(s, 2.6, 2.9, 2.15, 2.15, ORANGE, None)
text(s, 2.6, 2.9, 2.15, 2.15, [('25', 56, True, WHITE)], anchor=D.MSO_ANCHOR.MIDDLE, align=D.PP_ALIGN.CENTER)
text(s, 0.9, 5.35, 5.55, 0.6, [('番茄钟：一段专注 + 5 分钟休息', 15, False, GRAY)], align=D.PP_ALIGN.CENTER)
cards = [
    ('50 / 90 分钟', '写材料、改论文用长钟', BLUE),
    ('响了就歇', '休息不是偷懒，是给脑子存档', ORANGE),
    ('可叠模块', '在你工作台里就是一个计时按钮', TEAL),
]
for i, (t, d, c) in enumerate(cards):
    y = 2.1 + i * 1.55
    card(s, 6.9, y, 5.5, 1.35, WHITE, LINE)
    text(s, 7.25, y + 0.18, 4.9, 0.5, [(t, 18, True, INK)])
    text(s, 7.25, y + 0.66, 4.9, 0.6, [(d, 14, False, GRAY)])

# ============ P8 资料窗 + 只读铁律 ============
s = new_page('第一站 · 资料窗', '能翻你的资料，但不能改你的资料')
pic(s, os.path.join(FIGS, 'fig_readonly.png'), 0.7, 1.7, w=8.4)
demo_link(s, 0.9, 6.5)
text(s, 9.4, 1.9, 3.1, 4.3, [
    [('一句话守住安全：', 16, True, ORANGE)],
    ['把资料窗说成', 14.5, False, INK],
    [('「只能看、不能改', 16, True, INK),
     ('\n我的原文件」', 16, True, INK)],
    ('', 8, False, INK),
    ['AI 就只读取、不动你的源数据。', 14, False, GRAY],
], spacing=1.3, space_after=8)

# ============ P9 章节 02 ============
section_page('02', '跟做', '把需求说成大白话，工作台自己长出来')

# ============ P10 大白话范式 ============
s = new_page('第二站 · 一句话', '最小版，直接说给 AI 听')
card(s, 0.9, 1.95, 11.53, 2.5, WHITE, ORANGE, 1.6)
text(s, 1.3, 1.95, 10.7, 2.5, [
    [('「帮我做一个网页：', 21, True, INK), ('每天把要忙的事写上去，', 21, True, BLUE),
     ('按重要和着急分四块。', 21, True, BLUE)],
    [('这个网页要能保存——关了再打开，写的东西还在。', 21, True, INK),
     ('双击就能打开，不用装任何软件。」', 21, True, INK)],
], anchor=D.MSO_ANCHOR.MIDDLE, spacing=1.35, space_after=12)
card(s, 0.9, 4.7, 5.55, 1.55, ORANGE_SOFT, None)
text(s, 1.3, 4.7, 4.9, 1.55, [
    [('加一句更稳：', 15, True, ORANGE)],
    [('「内容存在我自己的电脑上，别存到网上。」', 16, True, ORANGE)],
], anchor=D.MSO_ANCHOR.MIDDLE, spacing=1.3)
card(s, 6.88, 4.7, 5.55, 1.55, WHITE, LINE)
text(s, 7.26, 4.7, 4.85, 1.55, [
    [('老师要做的只有一件：', 15, True, INK)],
    [('把「我想有什么」说清楚。', 17, True, INK)],
    [('不用写代码、不用装环境。', 14, False, GRAY)],
], anchor=D.MSO_ANCHOR.MIDDLE, spacing=1.3, space_after=5)
demo_link(s, 0.9, 6.5)

# ============ P11 章节 03 ============
section_page('03', '带走', '课下自己接着搭，才是真学会')

# ============ P12 迁移卡四步 ============
s = new_page('第三站 · 迁移', '回家也能自己搭：四步')
steps = [
    ('列', '写下这月最卡的 3 件事', ORANGE),
    ('配', '每件事对应一块拼图', BLUE),
    ('说', '用大白话说给 AI 听', TEAL),
    ('开加', '双击打开填自己的事，哪不对说哪', INK_SOFT),
]
y = 1.95
for i, (k, d, c) in enumerate(steps):
    card(s, 0.9, y, 11.53, 0.95, WHITE, LINE)
    rect(s, 0.9, y, 0.14, 0.95, c)
    oval(s, 1.3, y + 0.165, 0.62, 0.62, c, None)
    text(s, 1.3, y + 0.165, 0.62, 0.62, [(k, 14.5, True, WHITE)],
         anchor=D.MSO_ANCHOR.MIDDLE, align=D.PP_ALIGN.CENTER)
    text(s, 2.25, y, 9.9, 0.95, [(d, 17, True, INK)],
         anchor=D.MSO_ANCHOR.MIDDLE)
    y += 1.12
conclusion_bar(s, [('今天你做出一个；照这四步，你能做出十个。', 16.5, True, WHITE)], y=6.42, h=0.52)

# ============ P13 小结 ============
s = new_page('小结', '今天带走三句话')
rows = [
    ('①', '工作台 = 四块拼图里挑', '先挑最卡的 1~2 块，别贪全', ORANGE),
    ('②', '需求说成大白话', '「分四块 + 能保存 + 双击开」足矣', BLUE),
    ('③', '资料窗只读不改', '一句「不能改我的原文件」守住安全', TEAL),
]
for i, (num, t, d, c) in enumerate(rows):
    y = 1.98 + i * 1.42
    card(s, 0.9, y, 11.53, 1.2, WHITE, LINE)
    rect(s, 0.9, y, 0.14, 1.2, c)
    text(s, 1.35, y, 0.8, 1.2, [(num, 26, True, c)], anchor=D.MSO_ANCHOR.MIDDLE)
    text(s, 2.25, y, 4.3, 1.2, [(t, 19.5, True, INK)], anchor=D.MSO_ANCHOR.MIDDLE)
    text(s, 6.7, y, 5.4, 1.2, [(d, 15.5, False, GRAY)], anchor=D.MSO_ANCHOR.MIDDLE)

# ============ P14 结束页 ============
D._n[0] += 1
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, W, H, INK)
oval(s, -2.4, 3.4, 6.2, 6.2, None, INK_SOFT, 2.2)
oval(s, -1.4, 4.4, 4.2, 4.2, None, ORANGE, 1.6)
text(s, 0.95, 1.05, 8.0, 0.45, [('课后', 15, True, ORANGE)])
text(s, 0.95, 1.5, 9.0, 0.95, [('动手 & 下节预告', 36, True, WHITE)])
items = [
    ('必做', '用今天的一句话，生成一个你自己的最小版工作台'),
    ('动手', '打开 05_交互演示.html，点四步看它怎么搭起来'),
    ('下节', '给工作台加「AI 搭子」：把一周的事整理成给 AI 的指令'),
]
for i, (a, b) in enumerate(items):
    y = 2.95 + i * 1.18
    rect(s, 0.95, y, 1.05, 0.44, ORANGE if i == 0 else INK_SOFT, None,
         shape=D.MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, 0.95, y - 0.01, 1.05, 0.46, [(a, 13.5, True, WHITE)],
         anchor=D.MSO_ANCHOR.MIDDLE, align=D.PP_ALIGN.CENTER)
    text(s, 2.35, y - 0.12, 9.9, 0.7, [(b, 16.5, False, WHITE)],
         anchor=D.MSO_ANCHOR.MIDDLE, spacing=1.15)
text(s, 0.95, 6.6, 10.0, 0.4, [('下节课每人带自己的台子来，我们加模块。', 14, True, FAINT)])
D.footer_dark(s)

prs.save(OUT)
print('saved', OUT, len(prs.slides._sldIdLst), 'slides')
