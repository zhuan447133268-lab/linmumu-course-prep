# -*- coding: utf-8 -*-
"""02_课件.pptx —— 工科实例模板（机械臂正运动学 18 页）。
设计系统全部来自 deck_lib.py——本文件只含内容，不含排版。
换主题 = 复制本文件，改下方内容部分；课程名/配色/页数在 CONFIG 一行改。
用法：python make_pptx_v2.py [学科族] [输出pptx路径] [配图目录]
  默认：工科  ./02_课件.pptx  ./_figs
依赖：python-pptx；配图先由 make_figs.py 生成到配图目录。
"""
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os
import sys

# ---------- CONFIG（换主题只改这一行 + 内容部分） ----------
THEME  = sys.argv[1] if len(sys.argv) > 1 else "工科"
OUT    = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.path.join(os.getcwd(), "02_课件.pptx")
FIGS   = os.path.abspath(sys.argv[3]) if len(sys.argv) > 3 else os.path.join(os.getcwd(), "_figs")
COURSE_LABEL = "《工业机器人技术基础》 · 第 6 节 正运动学"
TOTAL_PAGES  = 18

# ---------- 设计系统（canonical 来源：deck_lib.py，勿在此复制） ----------
import deck_lib as D
D.new_prs(theme=THEME, course=COURSE_LABEL, total=TOTAL_PAGES)
from deck_lib import (INK, INK_SOFT, ORANGE, ORANGE_SOFT, TEAL, TEAL_SOFT,
                      BLUE, CREAM, WHITE, LINE, GRAY, FAINT, TRACK,
                      rect, oval, text, card, footer, footer_dark, header,
                      new_page, section_page, conclusion_bar, tag, demo_link)
prs = D.prs
W, H = D.W, D.H
_n = D._n
BLANK = prs.slide_layouts[6]

def pic(s, name, l, t, w=None, h=None):
    """本模板特有：按文件名从 FIGS 目录取图（deck_lib.pic 收全路径）。"""
    return D.pic(s, os.path.join(FIGS, name), l, t, w=w, h=h)

os.makedirs(os.path.dirname(OUT) or '.', exist_ok=True)

# P1 封面
# ============================================================
_n[0] += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, INK)
oval(s, 8.6, 1.2, 6.6, 6.6, None, INK_SOFT, 2.4)
oval(s, 9.7, 2.3, 4.4, 4.4, None, INK_SOFT, 2.0)
oval(s, 10.6, 3.2, 2.6, 2.6, None, ORANGE, 1.8)
oval(s, 11.62, 4.22, 0.56, 0.56, ORANGE)
text(s, 0.95, 1.72, 8.0, 0.4,
     [('工业机器人技术基础 · 第 6 节', 15, True, ORANGE)])
text(s, 0.95, 2.25, 9.2, 2.4, [
    ('平面两连杆机械臂', 46, True, WHITE),
    ('的正运动学', 46, True, WHITE),
], spacing=1.12)
rect(s, 0.98, 4.62, 1.55, 0.05, ORANGE)
text(s, 0.95, 4.88, 8.6, 0.5,
     [('给定关节角度，求出末端位置', 19, False, FAINT)])
text(s, 0.95, 6.35, 10.0, 0.4,
     [('工业机器人技术专业 · 大二 · 专业核心课 · 2 学时', 13.5, False,
       GRAY)])
footer_dark(s)

# ============================================================
# P2 今天的问题
# ============================================================
s = new_page('今天的问题', '大臂抬 50°、小臂再折 25°，末端落在哪儿？')
card(s, 0.9, 1.85, 5.9, 1.62, WHITE, LINE)
text(s, 1.3, 2.08, 5.2, 0.55, [
    [('θ₁ = 50°', 26, True, BLUE), ('　', 26, False, INK),
     ('θ₂ = 25°', 26, True, ORANGE)],
])
text(s, 1.3, 2.78, 5.2, 0.5, [
    [('→ (x, y) = ?', 26, True, INK)],
])
# 空间转换
text(s, 0.9, 3.95, 5.9, 1.6, [
    [('关节空间', 17, True, INK), ('　两个角度　→　', 17, False, GRAY),
     ('笛卡尔空间', 17, True, INK), ('　一个坐标', 17, False, GRAY)],
    ('这个方向叫正运动学；反方向叫逆运动学，下节课讲。', 14.5, False, GRAY),
], spacing=1.25, space_after=12)
card(s, 0.9, 5.05, 5.9, 0.92, TEAL_SOFT, None)
text(s, 1.3, 5.05, 5.2, 0.92,
     [('今天只解决这一个方向的转换。', 16, True, TEAL)],
     anchor=MSO_ANCHOR.MIDDLE)
pic(s, 'fig_arm.png', 7.05, 1.62, w=5.55)

# ============================================================
# P3 路线图
# ============================================================
s = new_page('本节路线', '三站，90 分钟')
rows = [
    ('01', '建系与难点', '建坐标系 · θ₂ 到底是谁和谁的夹角', BLUE),
    ('02', '公式与例题', '两段累加 · 特殊角与非特殊角各一道', ORANGE),
    ('03', '工作空间与应用', '圆环边界 · 上下料案例 · 量级校验', TEAL),
]
cw, gap, x0 = 3.72, 0.35, 0.9
for i, (num, t, d, c) in enumerate(rows):
    x = x0 + i * (cw + gap)
    card(s, x, 2.15, cw, 3.3, WHITE, LINE)
    rect(s, x, 2.15, cw, 0.14, c, None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.5)
    text(s, x + 0.35, 2.62, cw - 0.7, 0.85, [(num, 40, True, c)])
    text(s, x + 0.35, 3.62, cw - 0.7, 0.55, [(t, 21, True, INK)])
    text(s, x + 0.35, 4.28, cw - 0.7, 1.0, [(d, 14.5, False, GRAY)],
         spacing=1.25)

# ============================================================
# P4 章节 01
# ============================================================
section_page('01', '建系与难点', '先定基准，再攻下 θ₂ 这个每年必错的地方')

# ============================================================
# P5 坐标系与参数
# ============================================================
s = new_page('第一站 · 建系', '一切计算的前提：先建坐标系')
pic(s, 'fig_arm.png', 0.95, 1.72, w=5.25)
px = 7.35
items = [
    ('L₁ = 400 mm', '大臂长度', BLUE, WHITE),
    ('L₂ = 300 mm', '小臂长度', BLUE, WHITE),
    ('θ₁', '大臂与水平基准线的夹角', BLUE, WHITE),
]
for i, (a, b, c, bg) in enumerate(items):
    y = 1.85 + i * 1.06
    card(s, px, y, 5.05, 0.88, bg, LINE)
    text(s, px + 0.32, y, 4.5, 0.88, [
        [(a, 19, True, c), ('　' + b, 14.5, False, GRAY)]],
        anchor=MSO_ANCHOR.MIDDLE)
card(s, px, 1.85 + 3 * 1.06, 5.05, 0.88, ORANGE_SOFT, None)
text(s, px + 0.32, 1.85 + 3 * 1.06, 4.5, 0.88, [
    [('θ₂', 19, True, ORANGE), ('　待定——下一页重点', 14.5, True, ORANGE)]],
    anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.95, 6.5, 5.6, 0.45,
     [('基座为原点 O，水平向右为 x，竖直向上为 y。', 14.5, True, GRAY)])

# ============================================================
# P6 θ₂ 难点
# ============================================================
s = new_page('第一站 · 难点', 'θ₂ 到底是谁和谁的夹角？')
card(s, 0.9, 1.9, 5.75, 2.5, WHITE, LINE)
text(s, 1.28, 2.2, 5.0, 2.0, [
    [('✔ ', 20, True, TEAL), ('连杆 2 相对 连杆 1 的夹角', 19, True, INK)],
    [('✘ ', 20, True, ORANGE), ('不是连杆 2 相对水平线的角度', 19, False,
                                GRAY)],
], spacing=1.3, space_after=16)
card(s, 0.9, 4.62, 5.75, 1.28, TEAL_SOFT, None)
text(s, 1.28, 4.62, 5.0, 1.28, [
    [('记牢这一句：', 14.5, True, TEAL)],
    [('θ₂ 永远是相对角，求末端方向必须累加。', 18.5, True,
      TEAL)],
], anchor=MSO_ANCHOR.MIDDLE, spacing=1.25, space_after=8)
pic(s, 'fig_theta2.png', 6.95, 1.62, w=5.7)
demo_link(s, 0.95, 6.55)

# ============================================================
# P7 代价：大数字
# ============================================================
s = new_page('第一站 · 难点', '用错一次，代价是多少？')
text(s, 0.9, 2.5, 11.53, 0.5,
     [('同样给 θ₁ = 50°、θ₂ = 25°，两种理解，末端差出去多远：', 16.5, False,
       GRAY)], align=PP_ALIGN.CENTER)
text(s, 0.9, 3.05, 11.53, 1.9,
     [('253.6 mm', 92, True, ORANGE)], align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
conclusion_bar(s, [('两根杆长度加起来才 700 mm——末端偏出去 253.6 mm，工件直接抓空。',
                    16.5, True, WHITE)], y=5.35, h=0.66)
demo_link(s, 0.95, 6.4)

# ============================================================
# P8 章节 02
# ============================================================
section_page('02', '公式与例题', '两段累加，一道特殊角、一道真实角')

# ============================================================
# P9 第一步
# ============================================================
s = new_page('第二站 · 公式', '第一步：先算肘关节在哪')
card(s, 0.9, 2.0, 5.55, 2.15, WHITE, LINE)
text(s, 1.3, 2.0, 4.8, 2.15, [
    [('x₁ = L₁ · cos θ₁', 27, True, BLUE)],
    [('y₁ = L₁ · sin θ₁', 27, True, BLUE)],
], anchor=MSO_ANCHOR.MIDDLE, spacing=1.35, space_after=10)
card(s, 0.9, 4.42, 5.55, 1.5, TEAL_SOFT, None)
text(s, 1.3, 4.42, 4.85, 1.5, [
    [('就是直角三角形：', 15.5, True, TEAL)],
    [('水平方向配 cos，竖直方向配 sin。', 15.5, False,
      TEAL)],
], anchor=MSO_ANCHOR.MIDDLE, spacing=1.3)
pic(s, 'fig_accum.png', 6.75, 1.7, w=5.85)

# ============================================================
# P10 第二步：完整公式
# ============================================================
s = new_page('第二站 · 公式', '第二步：从肘关节再走一段')
card(s, 0.9, 1.98, 11.53, 2.05, WHITE, ORANGE, 1.6)
text(s, 1.3, 1.98, 10.7, 2.05, [
    [('x = L₁·cos θ₁ + ', 26, True, BLUE), ('L₂·cos(θ₁+θ₂)', 26, True,
                                             ORANGE)],
    [('y = L₁·sin θ₁ + ', 26, True, BLUE), ('L₂·sin(θ₁+θ₂)', 26, True,
                                             ORANGE)],
], anchor=MSO_ANCHOR.MIDDLE, spacing=1.35, space_after=10)
card(s, 0.9, 4.35, 5.55, 1.55, ORANGE_SOFT, None)
text(s, 1.3, 4.35, 4.9, 1.55, [
    [('最容易错的一步：', 15, True, ORANGE)],
    [('第二段的方向是 θ₁+θ₂，不是 θ₂。', 17, True, ORANGE)],
], anchor=MSO_ANCHOR.MIDDLE, spacing=1.3, space_after=6)
card(s, 6.88, 4.35, 5.55, 1.55, WHITE, LINE)
text(s, 7.26, 4.35, 4.85, 1.55, [
    [('结构只有一句话：', 15, True, INK)],
    [('每段 = 长度 × 方向余弦，从基座逐段累加。', 15.5, False, INK)],
    [('三连杆？再加一项。', 13.5, False, GRAY)],
], anchor=MSO_ANCHOR.MIDDLE, spacing=1.25, space_after=5)

# ============================================================
# P11 例题 1
# ============================================================
s = new_page('第二站 · 例题 1', '特殊角：先撑住信心')
tag(s, 9.35, 0.86, 'θ₁ = 60°　θ₂ = 30°', TEAL_SOFT,
    TEAL, w=3.05)
card(s, 0.9, 1.95, 5.55, 3.0, WHITE, LINE)
text(s, 1.28, 2.2, 4.9, 0.4, [('STEP 1 · 算肘关节', 14, True, BLUE)])
text(s, 1.28, 2.72, 4.9, 2.0, [
    ('x₁ = 400 × cos 60° = 200', 18, False, INK),
    ('y₁ = 400 × sin 60° = 346.4', 18, False, INK),
], spacing=1.3, space_after=10)
card(s, 6.88, 1.95, 5.55, 3.0, WHITE, LINE)
text(s, 7.26, 2.2, 4.9, 0.4, [('STEP 2 · 算末端', 14, True, ORANGE)])
text(s, 7.26, 2.72, 4.9, 2.0, [
    ('θ₁+θ₂ = 90°，cos 90° = 0，sin 90° = 1', 18, False, INK),
    ('x = 200 + 300 × 0 = 200.0', 18, False, INK),
    ('y = 346.4 + 300 × 1 = 646.4', 18, False, INK),
], spacing=1.3, space_after=10)
conclusion_bar(s, [('末端 ', 16.5, True, WHITE), ('(200.0, 646.4)', 20, True,
                   TEAL_SOFT)], y=5.35, h=0.66)

# ============================================================
# P12 例题 2
# ============================================================
s = new_page('第二站 · 例题 2', '非特殊角：考试的真实样子')
tag(s, 9.35, 0.86, 'θ₁ = 50°　θ₂ = 25°', ORANGE_SOFT, ORANGE, w=3.05)
card(s, 0.9, 1.95, 5.55, 3.0, WHITE, LINE)
text(s, 1.28, 2.2, 4.9, 0.4, [('STEP 1 · 查表 / 计算器', 14, True, BLUE)])
text(s, 1.28, 2.72, 4.9, 2.0, [
    ('cos 50° ≈ 0.643，sin 50° ≈ 0.766', 17, False, INK),
    ('θ₁+θ₂ = 75°，cos 75° ≈ 0.259', 17, False, INK),
    ('sin 75° ≈ 0.966', 17, False, INK),
], spacing=1.3, space_after=8)
card(s, 6.88, 1.95, 5.55, 3.0, WHITE, LINE)
text(s, 7.26, 2.2, 4.9, 0.4, [('STEP 2 · 代入累加', 14, True, ORANGE)])
text(s, 7.26, 2.72, 4.9, 2.0, [
    ('x = 257.1 + 77.6 = 334.7', 18, False, INK),
    ('y = 306.4 + 289.8 = 596.2', 18, False, INK),
], spacing=1.3, space_after=10)
conclusion_bar(s, [('两项都实实在在参与了计算——考试不会每次都给你凑好的角。',
                    15.5, True, WHITE)], y=5.35, h=0.66)

# ============================================================
# P13 章节 03
# ============================================================
section_page('03', '工作空间与应用', '够得着不只是一个感觉，是一条边界')

# ============================================================
# P14 工作空间
# ============================================================
s = new_page('第三站 · 工作空间', '是圆环，不是圆')
pic(s, 'fig_workspace.png', 1.05, 1.62, w=5.2)
px = 7.35
card(s, px, 1.85, 5.05, 1.28, WHITE, LINE)
text(s, px + 0.35, 1.85, 4.4, 1.28, [
    [('外半径　', 15.5, False, GRAY), ('L₁ + L₂ = 700 mm', 21, True, TEAL)],
], anchor=MSO_ANCHOR.MIDDLE)
card(s, px, 3.35, 5.05, 1.28, WHITE, LINE)
text(s, px + 0.35, 3.35, 4.4, 1.28, [
    [('内半径　', 15.5, False, GRAY), ('|L₁ − L₂| = 100 mm', 21, True,
                                       ORANGE)],
], anchor=MSO_ANCHOR.MIDDLE)
text(s, px + 0.05, 4.95, 5.0, 1.1, [
    ('太远够不着（超出 700）', 15.5, False, INK),
    ('太近也够不着——两根杆叠不拢', 15.5, False, INK),
], spacing=1.3, space_after=8)
conclusion_bar(s, [('够得着 ⇔ 距离同时满足：大于 100 且 小于 700。', 16, True,
                    WHITE)], y=6.32, h=0.56)

# ============================================================
# P15 工程案例
# ============================================================
s = new_page('第三站 · 应用', '上下料工位：这两个点够得着吗？')
pic(s, 'fig_workspace_pts.png', 1.05, 1.62, w=5.2)
px = 7.35
card(s, px, 1.85, 5.05, 1.05, WHITE, LINE)
text(s, px + 0.35, 1.85, 4.4, 1.05, [
    [('A 传送带　', 15.5, False, GRAY), ('(500, 200)', 18, True, INK),
     ('　538 mm', 15.5, True, TEAL)],
], anchor=MSO_ANCHOR.MIDDLE)
card(s, px, 3.12, 5.05, 1.05, WHITE, LINE)
text(s, px + 0.35, 3.12, 4.4, 1.05, [
    [('B 机床卡盘 ', 15.5, False, GRAY), ('(−100, 550)', 18, True, INK),
     ('　559 mm', 15.5, True, TEAL)],
], anchor=MSO_ANCHOR.MIDDLE)
card(s, px, 4.42, 5.05, 0.95, TEAL_SOFT, None)
text(s, px + 0.35, 4.42, 4.4, 0.95,
     [('两点都在 100 ~ 700 之间——都够得着。', 16, True,
       TEAL)], anchor=MSO_ANCHOR.MIDDLE)
text(s, px + 0.05, 5.62, 5.0, 0.5,
     [('实际工程还要看障碍物与姿态，下节课讲。', 13.5, False, GRAY)])

# ============================================================
# P16 量级校验
# ============================================================
s = new_page('第三站 · 工程习惯', '一个比公式更值钱的习惯：算完先扫一眼')
cards = [
    ('700', '上限：超出就错', TEAL, WHITE),
    ('100', '下限：落进就错', ORANGE, WHITE),
    ('820 ✗', '一眼假——不用看过程', ORANGE, ORANGE_SOFT),
]
cw, gap, x0 = 3.72, 0.35, 0.9
for i, (big, d, c, bg) in enumerate(cards):
    x = x0 + i * (cw + gap)
    card(s, x, 2.2, cw, 2.9, bg, LINE if bg == WHITE else None)
    text(s, x, 2.62, cw, 1.15, [(big, 46, True, c)], align=PP_ALIGN.CENTER)
    text(s, x + 0.3, 4.05, cw - 0.6, 0.8, [(d, 15, False, INK)],
         align=PP_ALIGN.CENTER, spacing=1.2)
conclusion_bar(s, [('上班以后没人帮你检查过程——但你自己 10 秒就能量级校验。',
                    16, True, WHITE)], y=5.55, h=0.62)

# ============================================================
# P17 小结
# ============================================================
s = new_page('小结', '今天带走三句话')
rows = [
    ('①', 'θ₂ 是相对角', '求末端方向必须累加：θ₁ + θ₂', ORANGE),
    ('②', '正运动学 = 两段累加', '每段「长度 × 方向余弦」，逐段相加', BLUE),
    ('③', '工作空间是圆环', '100 ~ 700，先量级校验再谈过程', TEAL),
]
for i, (num, t, d, c) in enumerate(rows):
    y = 1.98 + i * 1.42
    card(s, 0.9, y, 11.53, 1.2, WHITE, LINE)
    rect(s, 0.9, y, 0.14, 1.2, c)
    text(s, 1.35, y, 0.8, 1.2, [(num, 26, True, c)],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.25, y, 4.3, 1.2, [(t, 19.5, True, INK)],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.7, y, 5.4, 1.2, [(d, 15.5, False, GRAY)],
         anchor=MSO_ANCHOR.MIDDLE)

# ============================================================
# P18 作业（结束页）
# ============================================================
_n[0] += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, INK)
oval(s, -2.4, 3.4, 6.2, 6.2, None, INK_SOFT, 2.2)
oval(s, -1.4, 4.4, 4.2, 4.2, None, ORANGE, 1.6)
text(s, 0.95, 1.05, 8.0, 0.45, [('课后', 15, True, ORANGE)])
text(s, 0.95, 1.5, 8.0, 0.95, [('作业与下节预告', 36, True, WHITE)])
items = [
    ('必做', '习题 6-3、6-4：完整写出两段累加过程'),
    ('动手', '打开 05_交互演示.html，把末端拖到 (400, 400)，试试用了几组角度'),
    ('下节', '逆运动学：给坐标反求角度——今天埋的伏笔到时候全用上'),
]
for i, (a, b) in enumerate(items):
    y = 2.95 + i * 1.18
    rect(s, 0.95, y, 1.05, 0.44, ORANGE if i == 0 else INK_SOFT, None,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, 0.95, y - 0.01, 1.05, 0.46, [(a, 13.5, True, WHITE)],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    text(s, 2.35, y - 0.12, 9.9, 0.7, [(b, 16.5, False, WHITE)],
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
text(s, 0.95, 6.6, 10.0, 0.4,
     [('下周抽查：请一位同学上台，把末端拖到指定坐标。', 14, True, FAINT)])
footer_dark(s)

prs.save(OUT)
print('saved', OUT, len(prs.slides._sldIdLst), 'slides')
