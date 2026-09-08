# -*- coding: utf-8 -*-
"""03 讲稿 / 04 测验 的 Word 输出设计系统（linmumu-course-prep 用）。

老师拿到的是 .docx：双击就开、能打印、能拿笔改。版式 helper 全部集中在
本文件，换主题/换配色只走 new_doc(theme=...)，不许在内容脚本里手写 python-docx 版式。

配色 canonical 来源仍是 deck_lib.PALETTES——本文件只引用，不复制色值。

用法（讲稿）：
    import docx_lib as X
    doc = X.new_doc('《自动控制原理》· 第 6 节 根轨迹')
    X.h1(doc, '一、导入（5-10′）')
    X.para(doc, '同学们，今天先看一个现象……')
    X.mark(doc, '提问：根轨迹从哪儿出发？')      # 交互标记行，主色加粗
    X.mark(doc, '停顿')

用法（测验）：
    doc = X.new_doc('《自动控制原理》· 第 6 节 根轨迹 · 课堂测验')
    X.q(doc, 'Q1（单选）根轨迹起始于', ['A. 开环极点', 'B. 开环零点'])
    X.answers_start(doc)                          # 分页 + 「答案与解析」标题
    X.answer(doc, 'Q1：A。根轨迹起于开环极点，终于开环零点。对应目标 1。')
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deck_lib import PALETTES  # 配色唯一来源

FONT = 'Microsoft YaHei'
_theme = PALETTES['工科']
DARK = RGBColor((_theme['dark'] >> 16) & 0xFF, (_theme['dark'] >> 8) & 0xFF,
                _theme['dark'] & 0xFF)
PRIMARY = RGBColor((_theme['primary'] >> 16) & 0xFF,
                   (_theme['primary'] >> 8) & 0xFF, _theme['primary'] & 0xFF)
GRAY = RGBColor(0x6B, 0x72, 0x80)


def new_doc(course='', theme='工科'):
    """新建讲稿/测验文档：雅黑字体链 + 学科配色标题。返回 doc。"""
    global DARK, PRIMARY
    p = PALETTES.get(theme, PALETTES['工科'])
    DARK = RGBColor((p['dark'] >> 16) & 0xFF, (p['dark'] >> 8) & 0xFF,
                    p['dark'] & 0xFF)
    PRIMARY = RGBColor((p['primary'] >> 16) & 0xFF, (p['primary'] >> 8) & 0xFF,
                       p['primary'] & 0xFF)

    doc = Document()
    for st_name, color, size in (('Normal', None, 11),
                                 ('Heading 1', DARK, 16),
                                 ('Heading 2', PRIMARY, 13)):
        st = doc.styles[st_name]
        st.font.name = FONT
        st.font.size = Pt(size)
        if color is not None:
            st.font.color.rgb = color
        st.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn('w:eastAsia'), FONT)
    sec = doc.sections[0]
    for m in ('top_margin', 'bottom_margin'):
        setattr(sec, m, Pt(54))
    if course:
        h = doc.add_heading(course, level=1)
        for r in h.runs:
            r.font.size = Pt(15)
    return doc


def _ea(run):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),
                                                          FONT)


def h1(doc, text):
    h = doc.add_heading(text, level=1)
    for r in h.runs:
        r.font.size = Pt(16)
        _ea(r)
    return h


def h2(doc, text):
    h = doc.add_heading(text, level=2)
    for r in h.runs:
        r.font.size = Pt(13)
        _ea(r)
    return h


def para(doc, text, bold=False, size=None, color=None):
    """正文一段。bold/size/color 缺省走 Normal（11pt 雅黑深灰黑）。"""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    _ea(r)
    return p


def mark(doc, text):
    """讲稿交互标记行：【提问：……】/【停顿】，主色加粗，一眼可扫。"""
    p = doc.add_paragraph()
    r = p.add_run('【' + text + '】')
    r.bold = True
    r.font.color.rgb = PRIMARY
    _ea(r)
    return p


def q(doc, stem, options=None):
    """测验一道题（学生区）：题干加粗题号，选项逐行缩进。"""
    p = doc.add_paragraph()
    r = p.add_run(stem)
    r.bold = True
    _ea(r)
    if options:
        for opt in options:
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Pt(18)
            orun = op.add_run(opt)
            _ea(orun)
    return p


def answers_start(doc):
    """答案区：分页符 + 标题 + 一句使用说明。只打印分页前内容即学生版。"""
    doc.add_page_break()
    h1(doc, '答案与解析')
    return para(doc, '（教师版内容：以下答案与解析不印发给学生。打印时只打印本页之前的内容即为学生版。）',
                size=9, color=GRAY)


def answer(doc, text):
    """答案区一条：题号 + 答案 + 解析 + 对应目标。"""
    p = doc.add_paragraph()
    r = p.add_run(text)
    _ea(r)
    return p


def save(doc, path):
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc.save(path)
    print('saved', path)
